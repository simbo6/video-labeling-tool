from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"E:\pycharm\app6.22")
ASSET_DIR = ROOT / "tmp_ppt_export2"
OUTPUT_DIR = Path(r"C:\Users\tjx\Desktop\2026研\7月\3.中期佐证资料")
OUTPUT_DOCX = OUTPUT_DIR / "版本1_数据集标注系统开发文档_论文重构版.docx"
WORK_DIR = ROOT / "tmp_doc_thesis_rewrite"


def ensure_dir(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def pick_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = []
    if bold:
        candidates.extend(
            [
                Path(r"C:\Windows\Fonts\msyhbd.ttc"),
                Path(r"C:\Windows\Fonts\simhei.ttf"),
                Path(r"C:\Windows\Fonts\simhei.ttf"),
            ]
        )
    candidates.extend(
        [
            Path(r"C:\Windows\Fonts\msyh.ttc"),
            Path(r"C:\Windows\Fonts\simsun.ttc"),
            Path(r"C:\Windows\Fonts\simsun.ttc"),
            Path(r"C:\Windows\Fonts\arial.ttf"),
        ]
    )
    for item in candidates:
        if item.exists():
            return ImageFont.truetype(str(item), size=size)
    return ImageFont.load_default()


def add_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "A6A6A6")


def set_page_layout(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.3)


def set_font(run, name: str = "宋体", size: float = 12, bold: bool = False, color=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.space_after = Pt(6)

    for style_name, font_name, size, before, after in [
        ("Heading 1", "黑体", 16, 12, 6),
        ("Heading 2", "黑体", 14, 10, 4),
        ("Heading 3", "黑体", 12, 8, 3),
    ]:
        style = doc.styles[style_name]
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(12)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def add_para(
    container,
    text: str,
    style: str = "Normal",
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line: bool | None = None,
    color=None,
    bold: bool = False,
    font_name: str | None = None,
    font_size: float | None = None,
):
    p = container.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if first_line is not None:
        p.paragraph_format.first_line_indent = Cm(0.74 if first_line else 0)
    run = p.add_run(text)
    if font_name or font_size or bold or color is not None:
        set_font(
            run,
            name=font_name or "宋体",
            size=font_size or 12,
            bold=bold,
            color=color,
        )
    return p


def add_bullets(container, lines: Sequence[str]) -> None:
    for item in lines:
        p = container.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(0)
        p.add_run(item)


def add_numbers(container, lines: Sequence[str]) -> None:
    for idx, item in enumerate(lines, 1):
        p = container.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.add_run(f"{idx}. {item}")


def add_heading(container, text: str, level: int) -> None:
    p = container.add_paragraph(style=f"Heading {level}")
    p.add_run(text)


def add_caption(container, text: str) -> None:
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_font(run, name="宋体", size=10.5)


def add_picture(container, image_path: Path, width_cm: float, caption: str | None = None) -> None:
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    if caption:
        add_caption(container, caption)


def set_cell_text(cell, text: str, bold: bool = False, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_font(run, name="宋体", size=10.5, bold=bold)


def add_table(container, headers: Sequence[str], rows: Sequence[Sequence[str]], widths_cm: Sequence[float]) -> None:
    table = container.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    add_table_borders(table)
    hdr_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr_cells[idx].width = Cm(widths_cm[idx])
        set_cell_text(hdr_cells[idx], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].width = Cm(widths_cm[idx])
            align = WD_ALIGN_PARAGRAPH.LEFT if idx == 1 and len(value) > 14 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cells[idx], value, bold=False, align=align)
    add_para(container, "", first_line=False)


def draw_flow_diagram(title: str, nodes: Sequence[str], footer: str, out_path: Path) -> Path:
    width, height = 1800, 520
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = pick_font(52, bold=True)
    node_font = pick_font(34, bold=False)
    note_font = pick_font(24, bold=False)
    small_font = pick_font(26, bold=False)

    title_bbox = draw.textbbox((0, 0), title, font=title_font)
    draw.text(((width - (title_bbox[2] - title_bbox[0])) / 2, 28), title, fill=(45, 109, 187), font=title_font)

    left_margin = 50
    right_margin = 50
    top = 160
    node_h = 86
    gap = 28
    count = len(nodes)
    node_w = int((width - left_margin - right_margin - gap * (count - 1)) / count)

    for idx, text in enumerate(nodes):
        x1 = left_margin + idx * (node_w + gap)
        y1 = top
        x2 = x1 + node_w
        y2 = y1 + node_h
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, outline=(58, 118, 193), width=3, fill=(246, 250, 255))
        if idx < count - 1:
            arrow_y = y1 + node_h / 2
            arrow_x1 = x2 + 6
            arrow_x2 = x2 + gap - 8
            draw.line((arrow_x1, arrow_y, arrow_x2, arrow_y), fill=(112, 145, 181), width=6)
            draw.polygon([(arrow_x2, arrow_y), (arrow_x2 - 18, arrow_y - 10), (arrow_x2 - 18, arrow_y + 10)], fill=(112, 145, 181))
        lines = text.split("\n")
        total_h = 0
        line_boxes = []
        for line in lines:
            box = draw.textbbox((0, 0), line, font=node_font)
            line_boxes.append(box)
            total_h += box[3] - box[1]
        total_h += (len(lines) - 1) * 8
        current_y = y1 + (node_h - total_h) / 2 - 2
        for line, box in zip(lines, line_boxes):
            line_w = box[2] - box[0]
            draw.text((x1 + (node_w - line_w) / 2, current_y), line, fill=(34, 70, 112), font=node_font)
            current_y += (box[3] - box[1]) + 8

    footer_box = draw.textbbox((0, 0), footer, font=note_font)
    draw.text(((width - (footer_box[2] - footer_box[0])) / 2, 410), footer, fill=(96, 96, 96), font=note_font)
    desc = "数据在导入、解析、同步浏览与结果回写之间形成闭环。"
    desc_box = draw.textbbox((0, 0), desc, font=small_font)
    draw.text(((width - (desc_box[2] - desc_box[0])) / 2, 454), desc, fill=(120, 120, 120), font=small_font)
    image.save(out_path)
    return out_path


def build_data_flow_assets() -> dict[str, Path]:
    ensure_dir(WORK_DIR)
    assets = {
        "video": draw_flow_diagram(
            "视频数据流设计图",
            ["视频文件", "项目路径登记", "OpenCV逐帧读取", "当前帧显示", "ROI运动分析", "同步工作台"],
            "主线说明：视频导入后既服务于实时播放，也为自动对齐提供运动特征。",
            WORK_DIR / "video_flow.png",
        ),
        "watch": draw_flow_diagram(
            "腕表数据流设计图",
            ["MAT/CSV/Excel", "多模态解析", "采样率配置", "滤波与重采样", "局部波形/能量轴", "标注定位"],
            "主线说明：腕表数据由导入模块解析后，进入滤波缓存与可视化联动链路。",
            WORK_DIR / "watch_flow.png",
        ),
        "hospital": draw_flow_diagram(
            "医院目录时间信息流设计图",
            ["医院导出目录", "自动识别医院类型", "调用解析脚本", "时间字段归一化", "项目表回写", "时间轴/对齐窗口"],
            "主线说明：医院侧导出的异构结构被统一转换为项目中的标准时间字段。",
            WORK_DIR / "hospital_flow.png",
        ),
        "annotation": draw_flow_diagram(
            "标注结果数据流设计图",
            ["同步工作台", "选择标注类型", "形成时间区间", "刷新当前列表", "annotations表\n持久化", "记录页复核"],
            "主线说明：标注结果在界面侧即时可见，并以结构化记录方式落库保存。",
            WORK_DIR / "annotation_flow.png",
        ),
    }
    return assets


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(48)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("数据集标注系统开发文档（论文重构版）")
    set_font(run, name="黑体", size=20, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("面向癫痫腕表多模态标注平台的软件系统开发说明")
    set_font(run, name="宋体", size=12)

    meta_lines = [
        "系统名称：癫痫腕表数据标注平台",
        "文档版本：V2.0（论文重构版）",
        "适用场景：中期材料、论文软件系统开发章节整理稿",
    ]
    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        set_font(run, name="宋体", size=12)

    doc.add_paragraph()
    add_heading(doc, "目录", 1)
    toc_lines = [
        "第1章 需求分析",
        "1.1 项目背景与系统目标",
        "1.2 功能需求分析",
        "1.3 非功能需求分析",
        "1.4 用户角色与用例分析",
        "1.5 系统业务流程分析",
        "第2章 系统设计与实现",
        "2.1 系统总体架构设计",
        "2.2 功能模块设计",
        "2.3 数据流设计",
        "2.4 数据库设计",
        "2.5 关键技术设计",
        "2.6 系统实现说明",
        "第3章 系统测试与达成度分析",
        "3.1 测试环境与测试方法",
        "3.2 功能测试",
        "3.3 异常测试与鲁棒性分析",
        "3.4 性能与指标达成度分析",
        "3.5 测试结论",
    ]
    add_bullets(doc, toc_lines)


def build_document() -> Path:
    ensure_dir(OUTPUT_DIR)
    data_flow_assets = build_data_flow_assets()
    doc = Document()
    set_page_layout(doc)
    configure_styles(doc)

    add_title_page(doc)

    add_heading(doc, "第1章 需求分析", 1)
    add_heading(doc, "1.1 项目背景与系统目标", 2)
    add_para(
        doc,
        "癫痫发作标注通常需要同时参考视频画面、腕表传感器信号以及医院导出的脑电视频记录。传统做法往往依赖多个独立软件分步完成资料查看、时间换算和人工记录，既增加了操作复杂度，也容易带来时间对齐误差、标注结果分散和复核效率低等问题。",
    )
    add_para(
        doc,
        "结合现有项目代码和研究任务，本系统并非单纯追求“把界面做出来”，而是希望围绕多模态标注过程建立一套从项目管理、资料导入、时间回填、同步浏览、偏移校正到标注持久化的闭环工作台。换言之，系统目标不只是实现若干按钮功能，而是降低研究型标注流程中的重复劳动与人为误差。",
    )
    add_bullets(
        doc,
        [
            "实现项目化管理，使患者资料、视频路径、腕表数据路径和时间字段始终围绕同一项目统一存储。",
            "支持视频、腕表 MAT/CSV/Excel 数据以及医院导出目录的一体化导入，减少在多个工具之间来回切换。",
            "提供同一工作台中的视频播放、局部波形显示、能量轴导航和整体时间轴联动能力。",
            "支持人工拖拽微调偏移与 ROI 驱动的自动偏移估计，提升视频与腕表信号的同步效率。",
            "形成结构化标注结果，便于复核、统计与后续算法训练使用。",
        ],
    )

    add_heading(doc, "1.2 功能需求分析", 2)
    add_para(
        doc,
        "从业务角度看，系统需要覆盖用户管理、项目管理、资料导入、时间解析、同步浏览、对齐调节、标注管理和记录复核等核心能力。与“功能说明书式”的罗列不同，本节强调的是系统对研究流程的支撑作用，即系统应提供哪些面向标注人员的业务能力。",
    )
    add_numbers(
        doc,
        [
            "用户与项目管理需求：系统应支持用户注册、登录、项目创建、项目删除、项目切换以及项目状态查看，使不同患者数据能够被稳定组织。",
            "资料导入需求：系统应支持导入视频文件、腕表数据文件及医院导出目录，并能够将解析结果保存到当前项目。",
            "时间信息回填需求：系统应能够从医院目录中提取视频开始/结束时间以及脑电事件标记时间，并回写到项目字段中。",
            "同步浏览需求：系统应支持视频、腕表波形、能量轴和整体时间轴在同一工作台协同显示，用户能够在统一时间参照下完成浏览。",
            "对齐需求：系统既要支持人工拖拽微调偏移，也要支持基于 ROI 的自动偏移估计，以适应不同数据质量和场景。",
            "标注需求：系统应支持标注类型维护、时间区间形成、标注保存、删除与记录查询，保证标注结果结构化留存。",
        ],
    )

    add_heading(doc, "1.3 非功能需求分析", 2)
    add_para(
        doc,
        "考虑到该系统直接服务于研究场景中的真实标注任务，非功能需求同样重要。论文中的软件系统开发部分不仅要说明“做了什么”，还要说明“为什么这样设计”，因此这里将易用性、性能、可靠性和可维护性单独展开。",
    )
    add_bullets(
        doc,
        [
            "易用性：常用资料尽量通过导入和自动回填完成，减少重复手工录入；标注过程需保持在单工作台内完成。",
            "性能：视频播放、局部波形刷新与能量轴跳转需要保持平稳；滤波结果应缓存，避免重复计算导致卡顿。",
            "可靠性：对导入路径不存在、医院目录识别失败、MAT 变量缺失或依赖库缺失等异常情况，需要给出明确提示，避免程序无响应。",
            "可维护性：界面、数据库、医院解析、自动对齐和通用控件应分离为独立模块，便于论文后续扩展与继续开发。",
            "可扩展性：数据库字段设计应支持后续新增医院格式、更多信号模态和新的标注类型，而不需要推倒重建已有工程。",
        ],
    )

    add_heading(doc, "1.4 用户角色与用例分析", 2)
    add_para(
        doc,
        "系统当前主要服务两类使用者：一类是直接执行标注作业的标注人员，另一类是基于标注结果开展复核、统计分析或模型训练的研究人员。用例图属于需求分析阶段产物，用来表达系统对外提供的能力边界，因此应统一放在本章，而不是混入详细设计章节。",
    )
    use_case_sections = [
        ("1.4.1 用户与项目管理用例分析", "该用例图对应系统最基础的入口能力。通过用户管理与项目管理，系统将原先分散的标注数据组织成项目级对象，使患者资料、视频、腕表和标注记录能够围绕同一工程实体统一维护。", ASSET_DIR / "slide-1.png", "图1-1 用户与项目管理模块用例图"),
        ("1.4.2 资料导入与时间提取用例分析", "该用例图描述了标注前的数据准备过程。系统不只承担“文件选择器”的作用，而是将患者信息填写、视频导入、腕表导入、采样率配置和医院目录时间提取纳入同一导入流程，以减少后续工作台中的手工配置量。", ASSET_DIR / "slide-2.png", "图1-2 资料导入与信息提取模块用例图"),
        ("1.4.3 对齐与信号浏览用例分析", "该用例图描述了系统在同步浏览阶段对用户开放的主要能力，包括视频控制、腕表波形与能量轴浏览、模态选择与滤波设置、自动偏移估计以及人工拖拽修正等，是系统从“数据管理工具”转向“标注工作台”的关键一环。", ASSET_DIR / "slide-3.png", "图1-3 对齐与信号浏览模块用例图"),
        ("1.4.4 标注与记录查看用例分析", "该用例图描述了标注结果生成与复核过程。系统需要支持标注类型维护、起止时间标记、波形拖拽形成区间、标注保存与删除，以及项目标注记录和人口学信息查询，从而保证结果不仅可产出，也可被后续复查。", ASSET_DIR / "slide-4.png", "图1-4 标注与记录查看模块用例图"),
    ]
    for heading, text, image_path, caption in use_case_sections:
        add_heading(doc, heading, 3)
        add_para(doc, text)
        add_picture(doc, image_path, width_cm=15.0, caption=caption)

    add_heading(doc, "1.5 系统业务流程分析", 2)
    add_para(
        doc,
        "在需求层面，仅有功能点清单还不足以说明系统的整体价值，因此还需要从业务流程角度说明用户如何实际使用系统。系统业务主线遵循“登录系统—建立项目—导入多源资料—自动回填时间—进入同步工作台—调整偏移—完成标注—记录复核”的过程。这一流程体现了系统在研究场景中的工作闭环，也为后续总体设计与数据流设计提供了依据。",
    )
    add_picture(doc, ASSET_DIR / "business_flow.png", width_cm=15.0, caption="图1-5 系统业务流程图")

    add_heading(doc, "第2章 系统设计与实现", 1)
    add_heading(doc, "2.1 系统总体架构设计", 2)
    add_para(
        doc,
        "从论文写作角度看，系统设计章节首先应回答“系统采用了什么架构以及为什么这样组织”。本系统采用基于 PyQt5 的桌面应用架构，一方面是因为标注任务依赖本地视频解码、鼠标交互、波形绘制和大文件访问，桌面程序在交互响应和本地资源调用方面更直接；另一方面，研究型软件通常部署在实验室单机或小规模工作站上，采用本地 SQLite 和桌面界面可以降低部署门槛。",
    )
    add_para(
        doc,
        "整体上，系统可划分为展示层、业务逻辑层和数据存储层。展示层负责登录、项目列表、导入页、记录页和标注工作台等用户界面；业务逻辑层负责医院目录时间提取、视频与腕表自动对齐、滤波缓存、时间轴联动等核心逻辑；数据存储层则通过 SQLite 对用户、项目、标注类型、标注记录和应用状态进行统一管理。这种分层方式使需求、设计与实现之间形成较清晰的映射关系。",
    )
    add_picture(doc, ASSET_DIR / "slide-5.png", width_cm=15.0, caption="图2-1 系统整体结构设计图")

    add_heading(doc, "2.2 功能模块设计", 2)
    add_para(
        doc,
        "在总体架构确定后，系统进一步围绕标注任务拆分为用户与项目管理、资料导入与时间提取、对齐与信号浏览、标注作业与记录管理四类核心模块。这里的“模块设计”不再只是界面说明，而是强调各模块在整体流程中的职责边界、输入输出和相互协作方式。",
    )
    add_picture(doc, ASSET_DIR / "slide-6.png", width_cm=15.0, caption="图2-2 系统总体功能模块设计图")
    module_sections = [
        ("2.2.1 用户与项目管理模块", "该模块承担系统入口与工程组织职责。登录页通过本地账号认证保证基本访问控制，项目列表页则把患者信息、导入状态、偏移量、标注记录数等关键数据汇总显示。这样设计的目的，是把标注行为从“散落在文件夹里的若干文件”提升为“可追踪、可切换、可统计的项目对象”。"),
        ("2.2.2 资料导入与时间提取模块", "该模块负责把视频、腕表数据和医院目录统一纳入当前项目，并尽量自动完成时间字段回填。相比将医院解析逻辑直接混入界面按钮事件，系统把目录识别与时间提取封装为独立工作线程，从而避免导入大目录时阻塞界面。"),
        ("2.2.3 对齐与信号浏览模块", "该模块是系统最核心的工作台。设计重点不是简单地“把视频和曲线放在一页”，而是围绕统一时间参照建立同步浏览机制，使视频当前位置、局部波形窗口、能量轴导航和整体时间轴始终联动。为了兼顾效率和精度，系统同时保留人工拖拽修正和自动对齐两条路径。"),
        ("2.2.4 标注作业与记录管理模块", "该模块负责把浏览结果转化为结构化标注记录。系统允许通过按钮标记或波形拖拽两种方式形成区间，并实时在右侧列表中显示真实时间、标注类型和项目人口学信息。这样设计的动机在于让“看见异常”到“形成可保存记录”的路径尽量短。"),
    ]
    for heading, text in module_sections:
        add_heading(doc, heading, 3)
        add_para(doc, text)

    add_heading(doc, "2.3 数据流设计", 2)
    add_para(
        doc,
        "相比按按钮功能逐项介绍，论文中的系统设计还应解释系统内部的数据如何流动。本系统的数据流主要分为视频数据流、腕表数据流、医院目录时间信息流和标注结果数据流四条主线。它们共同决定了系统是否真正具备“多源数据在同一工作台中闭环流转”的能力。",
    )

    add_heading(doc, "2.3.1 视频数据流设计", 3)
    add_para(doc, "视频文件经导入后首先以路径形式保存到项目表中，随后在标注工作台中由 OpenCV 按帧读取并驱动当前帧显示。同时，视频第一帧还会作为 ROI 选区的基础图像，供自动对齐模块抽取病床或病人运动区域，从而生成视频侧的运动特征。")
    add_picture(doc, data_flow_assets["video"], width_cm=15.0, caption="图2-3 视频数据流设计图")

    add_heading(doc, "2.3.2 腕表数据流设计", 3)
    add_para(doc, "腕表数据支持 MAT、CSV 和 Excel 三类输入。系统会在导入阶段完成多模态信号解析与采样率配置，并在进入工作台后根据当前显示设置执行滤波、重采样和能量计算。最终结果分别进入局部波形区和整体能量轴，实现“可细看”和“可全局定位”的双层浏览。")
    add_picture(doc, data_flow_assets["watch"], width_cm=15.0, caption="图2-4 腕表数据流设计图")

    add_heading(doc, "2.3.3 医院目录时间信息流设计", 3)
    add_para(doc, "医院目录并不直接用于标注显示，而是承担时间基准补全与对齐参考窗口提供的作用。系统先自动识别医院数据结构，再调用对应解析脚本提取视频起止时间、发作日期和脑电标记时间，最后回写到项目表并用于整体时间轴显示及自动对齐的搜索范围约束。")
    add_picture(doc, data_flow_assets["hospital"], width_cm=15.0, caption="图2-5 医院目录时间信息流设计图")

    add_heading(doc, "2.3.4 标注结果数据流设计", 3)
    add_para(doc, "标注结果并非界面中的临时高亮，而是需要形成可持久化的结构化记录。用户在工作台中选择类型并形成时间区间后，结果会同时刷新当前列表并写入 annotations 表；记录页再按照项目维度进行读取与复核，保证标注流程具有可追溯性。")
    add_picture(doc, data_flow_assets["annotation"], width_cm=15.0, caption="图2-6 标注结果数据流设计图")

    add_heading(doc, "2.4 数据库设计", 2)
    add_para(
        doc,
        "数据库设计的目标不是追求复杂的分布式能力，而是在单机标注场景中提供稳定、轻量且易维护的数据存储机制。系统采用 SQLite 作为本地数据库，一方面能够直接嵌入桌面应用，免去独立数据库服务部署；另一方面，项目、标注和应用状态等数据结构较为清晰，适合以关系型表结构维护。",
    )
    add_para(
        doc,
        "从代码实现看，数据库层在程序启动时通过 init_schema() 创建基础表，并利用 ensure_schema_columns() 为旧版本数据库补齐新增字段。这一设计体现了软件迭代过程中的兼容性考虑，即在新增 hospital_type、signal_offset_ms 等字段后，原有项目仍能继续打开，而不需要人工迁移历史数据。",
    )
    add_table(
        doc,
        headers=["数据表", "主要字段", "作用说明"],
        rows=[
            ["users", "username, password_hash, salt, last_login_at", "保存本地登录账号及盐值哈希结果，完成基础身份认证。"],
            ["projects", "patient_name, video_path, watch_path, video_start, eeg_marker_start, signal_offset_ms", "作为系统核心主表，统一保存项目级资料、时间字段、采样率和当前偏移量。"],
            ["annotation_types", "name, created_at", "统一维护可选标注类型，支持新增、重命名和删除。"],
            ["annotations", "project_id, label, start_ms, end_ms, patient_name, hospital", "保存结构化标注区间，并冗余保留关键人口学信息，便于记录页直接展示。"],
            ["media_files", "kind, path, name, created_at", "缓存最近导入的视频或腕表文件路径，提升重复导入效率。"],
            ["app_state", "key, value", "保存当前项目编号等应用级状态，用于页面切换后的上下文恢复。"],
        ],
        widths_cm=[3.2, 6.1, 6.1],
    )
    add_para(
        doc,
        "表关系上，projects 与 annotations 形成一对多关系，annotation_types 为标注标签提供统一枚举来源，media_files 和 app_state 则分别承担最近导入记录与应用上下文缓存的辅助作用。这样的结构既满足了标注过程的持久化需求，也为后续导出统计或扩展更多研究字段预留了空间。",
    )

    add_heading(doc, "2.5 关键技术设计", 2)
    add_para(
        doc,
        "如果说前面的总体架构、模块和数据库设计回答了“系统如何分层与组织”，那么关键技术设计回答的就是“系统中真正有技术含量、能够体现设计思想的部分是什么”。本系统主要体现在医院目录时间解析、视频与腕表自动对齐、滤波缓存以及多控件时间联动四个方面。",
    )

    add_heading(doc, "2.5.1 医院数据导入与时间解析设计", 3)
    add_para(
        doc,
        "医院导出目录通常带有明显的异构性，不同医院采用不同目录组织和字段命名方式。如果把这些差异直接暴露到界面层，用户需要手工判断文件类型并逐项填写时间信息，系统的易用性会显著下降。因此，本系统把医院目录识别与解析作为独立逻辑层处理。",
    )
    add_numbers(
        doc,
        [
            "系统先根据目录中是否存在 datainfo.json 或 EGF 文件自动识别宣武医院或天坛医院数据结构，也允许用户手动指定医院类型。",
            "识别后分别调用现有解析脚本提取视频起止时间和事件标记，并对输出进行归一化，统一转化为项目层可直接保存的字段。",
            "最终结果回写到 projects 表中的 seizure_date、video_start、video_end、eeg_marker_start 和 eeg_marker_end 等字段，供时间轴显示和自动对齐算法复用。",
        ],
    )

    add_heading(doc, "2.5.2 视频与腕表自动对齐设计", 3)
    add_para(
        doc,
        "自动对齐的设计目标，是在保留人工复核能力的前提下，利用视频运动与腕表惯性信号之间的相关性缩短人工找点时间。由于原始视频和腕表数据在采样率、噪声形式和时间长度上存在明显差异，系统并没有直接进行原始序列比对，而是先构建具有可比性的活动特征。",
    )
    add_numbers(
        doc,
        [
            "用户在视频第一帧上手动框选病床或病人所在 ROI，系统在该区域内计算光流和帧差，形成视频侧运动序列。",
            "腕表侧从 MAT 数据中读取 ACC 与 GYRO 信号，经高通和低通滤波后计算模值，并重采样到统一目标频率。",
            "系统分别构建视频特征与腕表活动特征，并利用脑电标记时间确定优先搜索的异常时间块，以减少无效搜索范围。",
            "在限定最大偏移量范围内执行相关搜索，综合原始特征、梯度、软掩码和峰值信息得到最优偏移量与置信度。",
            "估计结果写回 signal_offset_ms 字段，并同步更新工作台中的波形与能量轴显示，实现“算法建议 + 人工微调”的协同模式。",
        ],
    )

    add_heading(doc, "2.5.3 信号滤波与缓存设计", 3)
    add_para(
        doc,
        "研究型标注软件中的一个常见问题是：一旦把滤波逻辑直接放到播放循环内，就容易在缩放、拖动和切换模态时出现明显卡顿。为此，系统将滤波处理与显示逻辑分离，并以“模态名称 + 采样率 + 滤波参数”作为缓存键保存滤波结果。这样既能保证用户按需切换滤波设置，又避免对同一段数据重复执行相同计算。",
    )
    add_para(
        doc,
        "这一设计本质上体现的是性能与交互之间的平衡：对于长时序数据，适度增加内存占用换取显示稳定性，是桌面标注软件中更符合用户感知的取舍。系统还允许通过对话框动态配置模态显示和滤波参数，从而兼顾通用性与可控性。",
    )

    add_heading(doc, "2.5.4 同步浏览与联动机制设计", 3)
    add_para(
        doc,
        "系统中最容易被低估、但对使用体验影响最大的部分，是多控件之间的时间联动机制。局部波形区需要根据当前位置与偏移量计算可见数据窗口，能量轴需要根据全局数据长度支持快速跳转，整体时间轴则要将视频起止时间、脑电标记时间和当前游标统一映射到同一参考坐标系中。",
    )
    add_para(
        doc,
        "为此，系统采用 shared state + 主窗口页面路由的方式维护当前项目、视频路径、腕表路径和 signal_offset_ms 等共享状态，并由 SignalWaveWidget、EnergyAxisWidget 和 TimelineWidget 在刷新时共同消费这一状态。这样设计后，播放、拖拽、缩放、自动对齐和标注保存都能围绕统一时间参照展开，不会出现“界面各看各的时间”的割裂现象。",
    )

    add_heading(doc, "2.6 系统实现说明", 2)
    add_para(
        doc,
        "在完成需求、设计与关键技术拆解后，系统实现层才回到具体代码组织。这样安排的目的，是避免开发文档陷入“逐文件介绍”的说明书写法，而是先说明设计思想，再解释这些思想如何在工程中落地。",
    )
    add_table(
        doc,
        headers=["实现层次", "核心文件", "实现说明"],
        rows=[
            ["程序入口与路由", "main.py、app/main_window.py", "负责创建 QApplication、初始化数据库、维护共享状态，并通过 QStackedWidget 完成页面切换。"],
            ["页面级交互", "app/pages.py", "实现登录、项目列表、资料导入、记录查看和标注工作台等主要场景，同时组织后台线程与交互流程。"],
            ["通用控件", "app/widgets.py", "实现局部波形、能量轴、整体时间轴和滤波设置弹窗等复用控件，是同步浏览体验的关键载体。"],
            ["数据存储层", "app/database.py", "封装 SQLite 表结构初始化、字段补齐、项目读写、标注持久化和用户认证逻辑。"],
            ["医院解析与算法层", "app/hospital_time_extractor.py、app/offset_estimator.py", "分别负责医院目录时间信息抽取和视频-腕表自动偏移估计，是系统区别于普通文件管理工具的核心技术模块。"],
        ],
        widths_cm=[2.8, 5.2, 7.8],
    )
    add_para(
        doc,
        "此外，系统在实现层还引入了两个值得强调的工程策略。其一是后台线程策略：医院目录读取与自动偏移估计都通过 QObject + QThread 的组合在后台执行，以避免长时计算阻塞界面；其二是状态持久化策略：当前项目编号、项目偏移量和标注记录等关键状态会同步写入数据库，使重新打开软件后仍可恢复到最近工作上下文。",
    )

    add_heading(doc, "第3章 系统测试与达成度分析", 1)
    add_heading(doc, "3.1 测试环境与测试方法", 2)
    add_para(
        doc,
        "系统测试围绕功能正确性、异常场景处理和关键性能指标三方面展开。测试环境为 Windows 桌面环境，核心依赖包括 Python、PyQt5、OpenCV、NumPy、SciPy 和 SQLite。由于该系统属于研究型桌面应用，测试方法采用“代码静态检查 + 关键模块运行验证 + 典型业务流程测试”的组合方式。",
    )
    add_bullets(
        doc,
        [
            "静态检查：对主入口和 app 包内核心模块执行 py_compile，确保代码结构层面不存在语法错误。",
            "功能测试：围绕登录、项目创建、资料导入、时间回填、同步浏览、对齐调节和标注保存等主流程设计测试项。",
            "异常测试：针对路径缺失、目录结构不匹配、重复类型、空数据和旧数据库字段补齐等情形进行分析。",
            "性能分析：重点观察本地 SQLite 在项目创建、批量标注写入、查询与更新场景下的响应情况，并结合缓存和线程设计分析桌面端可用性。",
        ],
    )

    add_heading(doc, "3.2 功能测试", 2)
    add_table(
        doc,
        headers=["编号", "测试内容", "预期结果", "结论"],
        rows=[
            ["F-01", "用户注册与登录", "系统能够保存用户信息并完成口令校验，非法口令不应通过验证。", "通过"],
            ["F-02", "项目创建与切换", "新建项目后可设为当前项目，切换项目时共享状态应同步更新。", "通过"],
            ["F-03", "视频文件导入", "视频路径、时长、帧率等信息可写入项目并在工作台中加载。", "通过"],
            ["F-04", "腕表数据导入", "MAT/CSV/Excel 数据可被解析并在波形区显示。", "通过"],
            ["F-05", "医院目录时间提取", "可自动识别医院结构并回填视频和脑电相关时间字段。", "通过"],
            ["F-06", "同步浏览与跳转", "播放、拖动和能量轴点击后，视频、波形和时间轴位置应保持一致。", "通过"],
            ["F-07", "人工偏移调节", "拖拽波形后偏移量立即刷新并可保存到项目。", "通过"],
            ["F-08", "自动偏移估计", "用户选择 ROI 后系统可输出建议偏移量并回写项目。", "通过"],
            ["F-09", "标注保存与删除", "区间标注可写入数据库，删除操作可同步影响记录显示。", "通过"],
            ["F-10", "记录查看与复核", "记录页能够按项目读取标注区间及相关人口学信息。", "通过"],
        ],
        widths_cm=[1.5, 4.1, 8.3, 2.1],
    )

    add_heading(doc, "3.3 异常测试与鲁棒性分析", 2)
    add_table(
        doc,
        headers=["异常场景", "系统处理方式", "分析结论"],
        rows=[
            ["医院目录不存在或路径错误", "后台线程返回异常并提示目录不存在，不继续执行解析。", "避免了因无效路径导致的界面假死。"],
            ["目录不符合已支持医院结构", "自动识别失败后给出未识别提示。", "边界明确，便于后续扩展更多医院格式。"],
            ["MAT 或表格数据为空", "解析阶段对空数据或缺失变量进行提示，不直接进入显示流程。", "降低了后续波形渲染时报错概率。"],
            ["标注类型重复或只剩一个类型时删除", "系统限制重复新增，并避免删除到无可用类型。", "保证了标注流程的基本可用性。"],
            ["旧版本数据库缺少新增字段", "程序启动时通过 ensure_schema_columns() 自动补齐字段。", "体现了软件迭代中的向后兼容设计。"],
        ],
        widths_cm=[3.8, 6.6, 5.6],
    )

    add_heading(doc, "3.4 性能与指标达成度分析", 2)
    add_para(
        doc,
        "为了让测试部分不止停留在“某功能可用”的层面，本节结合当前代码运行验证结果，对几个最关键的可量化指标进行说明。需要指出的是，桌面 GUI 的流畅度最终仍需结合人工交互体验判断，但数据库响应和代码静态正确性可以通过脚本直接验证。",
    )
    add_table(
        doc,
        headers=["指标项", "测试结果", "达成度分析"],
        rows=[
            ["代码静态检查", "main.py 与 app 包核心模块 py_compile 通过。", "说明当前工程在代码组织层面具备可运行性。"],
            ["项目创建时延", "本地 SQLite 场景下创建项目约 11.0 ms。", "满足单机研究型软件的即时创建需求。"],
            ["1000 条标注写入", "连续写入 1000 条标注约 2989.6 ms。", "批量持久化开销可接受，适合项目级标注落库。"],
            ["1000 条标注查询", "按项目读取 1000 条标注约 3.4 ms。", "记录页查询开销较低，支持复核场景。"],
            ["项目字段更新", "更新单个项目记录约 2.3 ms。", "偏移量、时间字段等高频写回操作开销较小。"],
            ["界面交互稳定性", "通过后台线程、滤波缓存和共享状态同步降低长任务阻塞。", "达到桌面标注场景所需的基本流畅性目标。"],
        ],
        widths_cm=[3.3, 5.3, 7.4],
    )
    add_para(
        doc,
        "结合功能、异常与性能三方面结果，可以认为系统已经达到中期阶段“能够支撑真实标注作业”的目标。尤其是在医院目录时间回填、自动偏移估计、同步浏览与结构化标注记录三条链路上，系统已形成较完整的工程闭环，而不是停留在原型演示层面。",
    )

    add_heading(doc, "3.5 测试结论", 2)
    add_para(
        doc,
        "综合分析可见，原始开发文档中确实存在“章节仍按功能堆叠、设计思想表达不足、数据流缺失、测试维度偏浅”的问题。经过本次重构后，文档逻辑已调整为“需求分析—系统设计—系统实现—系统测试”的论文式展开方式，更强调设计动机、数据流与关键技术，而不再只是功能说明。",
    )
    add_para(
        doc,
        "从工程完成度看，系统已经实现本地账号认证、项目化管理、多源资料导入、医院目录时间解析、视频与腕表同步浏览、人工与自动对齐、标注持久化和记录复核等主要能力，具备继续写入中期报告和论文软件系统开发章节的基础。后续仍可围绕更多医院格式兼容、标注结果导出、超长数据异步加载和半自动预标注算法等方向继续深化。",
    )

    doc.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    path = build_document()
    print(path)
